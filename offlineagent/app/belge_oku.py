"""Belgelerden düz metin çıkarır. Her belge (sayfa etiketi, metin) listesine dönüşür.

Taranmış PDF sayfaları ve resim dosyaları için `ocr` verilirse yazı görsel modelle okunur.
"""
import base64
import io
from pathlib import Path
from typing import Callable

METIN_UZANTILARI = {".txt", ".md", ".csv", ".json", ".log"}
RESIM_UZANTILARI = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}
DESTEKLENEN = METIN_UZANTILARI | RESIM_UZANTILARI | {".pdf", ".docx", ".xlsx"}

# Bir PDF sayfasında bundan az karakter varsa taranmış (resim) sayfa sayılır
OCR_ESIGI = 20

Ocr = Callable[[str], str]  # base64 PNG/JPEG alır, metin döndürür


def resmi_hazirla(resim, en_fazla: int = 1600) -> str:
    """PIL resmini küçültüp base64 JPEG'e çevirir."""
    resim = resim.convert("RGB")
    resim.thumbnail((en_fazla, en_fazla))
    tampon = io.BytesIO()
    resim.save(tampon, "JPEG", quality=90)
    return base64.b64encode(tampon.getvalue()).decode("ascii")


def _pdf(yol: Path, ocr: Ocr | None) -> list[tuple[str, str]]:
    from pypdf import PdfReader

    okuyucu = PdfReader(str(yol))
    sayfalar = [(f"s.{i + 1}", sayfa.extract_text() or "") for i, sayfa in enumerate(okuyucu.pages)]
    taranmis = [i for i, (_, metin) in enumerate(sayfalar) if len(metin.strip()) < OCR_ESIGI]
    if ocr and taranmis:
        import pypdfium2

        pdf = pypdfium2.PdfDocument(str(yol))
        try:
            for i in taranmis:
                resim = pdf[i].render(scale=2).to_pil()
                sayfalar[i] = (f"s.{i + 1} (OCR)", ocr(resmi_hazirla(resim)))
        finally:
            pdf.close()
    return sayfalar


def _resim(yol: Path, ocr: Ocr | None) -> list[tuple[str, str]]:
    if not ocr:
        return []
    from PIL import Image

    with Image.open(yol) as resim:
        return [("OCR", ocr(resmi_hazirla(resim)))]


def _docx(yol: Path) -> list[tuple[str, str]]:
    import docx

    belge = docx.Document(str(yol))
    satirlar = [p.text for p in belge.paragraphs]
    for tablo in belge.tables:
        for satir in tablo.rows:
            satirlar.append(" | ".join(h.text.strip() for h in satir.cells))
    return [("", "\n".join(satirlar))]


def _xlsx(yol: Path) -> list[tuple[str, str]]:
    from openpyxl import load_workbook

    kitap = load_workbook(str(yol), read_only=True, data_only=True)
    try:
        sonuc = []
        for sayfa in kitap.worksheets:
            satirlar = []
            for satir in sayfa.iter_rows(values_only=True):
                metin = " | ".join("" if h is None else str(h) for h in satir)
                if metin.strip(" |"):
                    satirlar.append(metin)
            sonuc.append((f"sayfa: {sayfa.title}", "\n".join(satirlar)))
        return sonuc
    finally:
        kitap.close()


def _metin(yol: Path) -> list[tuple[str, str]]:
    ham = yol.read_bytes()
    for kodlama in ("utf-8-sig", "cp1254"):
        try:
            return [("", ham.decode(kodlama))]
        except UnicodeDecodeError:
            continue
    return [("", ham.decode("latin-1"))]


def belgeyi_oku(yol: Path, ocr: Ocr | None = None) -> list[tuple[str, str]]:
    uzanti = yol.suffix.lower()
    if uzanti == ".pdf":
        return _pdf(yol, ocr)
    if uzanti in RESIM_UZANTILARI:
        return _resim(yol, ocr)
    if uzanti == ".docx":
        return _docx(yol)
    if uzanti == ".xlsx":
        return _xlsx(yol)
    return _metin(yol)
